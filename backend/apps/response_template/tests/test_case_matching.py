"""Phase 2 新增逻辑测试: 案例列识别 + 案例匹配排序。"""

import pytest
from docx import Document

from apps.response_template.services.filler import OoxmlFiller

filler = OoxmlFiller()


def make_table_doc(header: list[str], rows: list[list[str]]):
    doc = Document()
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    for j, h in enumerate(header):
        t.cell(0, j).text = h
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            t.cell(i + 1, j).text = v
    return doc, t


class TestCaseColumnDetection:
    def test_detect_columns(self):
        doc, t = make_table_doc(
            ["项目起止年月", "项目名称", "项目甲方名称", "证明人", "实施金额", "项目范围概述", "备注"],
            [["", "", "", "", "", "", ""]],
        )
        col_map = filler._detect_case_columns(t)
        assert col_map[0] == "period"
        assert col_map[1] == "project_name"
        assert col_map[2] == "client_name"
        assert col_map[3] == "client_contact"
        assert col_map[4] == "amount"
        assert col_map[5] == "scope"
        assert col_map[6] == "remark"

    def test_detect_partial_columns(self):
        doc, t = make_table_doc(["项目名称", "甲方名称"], [["", ""]])
        col_map = filler._detect_case_columns(t)
        assert col_map == {0: "project_name", 1: "client_name"}


class TestCaseMatching:
    @pytest.mark.django_db
    def test_match_cases_filters_by_company(self):
        from django.contrib.auth import get_user_model

        from apps.enterprise.models import CompanyCase, CompanyProfile

        user = get_user_model().objects.create_user(
            username="case_tester", password="x"
        )
        company = CompanyProfile.objects.create(
            name="测试企业", is_default=True, status="active", created_by=user
        )
        CompanyCase.objects.create(
            company=company, project_name="XX银行安全众测项目", client_name="XX银行",
            created_by=user,
        )
        CompanyCase.objects.create(
            company=company, project_name="YY银行评估项目", client_name="YY银行",
            created_by=user,
        )

        from types import SimpleNamespace

        block = SimpleNamespace(
            template=SimpleNamespace(project=SimpleNamespace(name="安全众测服务采购项目"))
        )
        cases = filler._match_cases(block, limit=5)
        assert len(cases) == 2
        # 相关度排序: 含"安全众测"关键词的排前面
        assert cases[0].project_name == "XX银行安全众测项目"

    @pytest.mark.django_db
    def test_match_cases_empty(self):
        from types import SimpleNamespace

        block = SimpleNamespace(
            template=SimpleNamespace(project=SimpleNamespace(name="测试项目"))
        )
        assert filler._match_cases(block, limit=3) == []
